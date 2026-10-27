"""Fill the SCOPE Project I report template with the ClarifyReview report.

The template is edited in place: the cover, declaration, certificate,
acknowledgements, contents, lists and abstract keep the template's own
paragraphs and formatting, and the chapters are written into its chapter
sections, so roman numbering for the front matter and arabic numbering from
Chapter 1 come from the template itself. Text and numbers come from
report_content.py, which reads paper/numbers.json.

Page numbers in the contents and the lists are found, not guessed: the report is
built once, exported to PDF through Microsoft Word, each heading and caption is
located in the PDF, and the report is built again with those numbers.

    .venv/bin/python docs/build_report.py --template "<template.docx>" \\
        --out ~/Downloads/ClarifyReview_Review3/ClarifyReview_Project_I_Report.docx
"""

from __future__ import annotations

import argparse
import copy
import re
import subprocess
import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

sys.path.insert(0, str(Path(__file__).resolve().parent))
import report_content as C  # noqa: E402

FONT = "Times New Roman"
TEXT_WIDTH_IN = 6.1


# ------------------------------------------------------------ run helpers

def _font(run, size=None, bold=None, italic=None, name=FONT):
    run.font.name = name
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_text(p: Paragraph, text: str, keep_highlight: bool = False) -> None:
    """Replace a template paragraph's text, keeping its first run's formatting."""
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    first = runs[0]
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    first.text = text
    if not keep_highlight:
        first.font.highlight_color = None


def set_runs(p: Paragraph, parts: list[tuple[str, bool]]) -> None:
    """Rewrite a paragraph as (text, highlighted) parts in its first run's format."""
    proto = copy.deepcopy(p.runs[0]._r)
    for r in p.runs:
        r._r.getparent().remove(r._r)
    for text, hl in parts:
        el = copy.deepcopy(proto)
        p._p.append(el)
        from docx.text.run import Run
        run = Run(el, p)
        run.text = text
        run.font.highlight_color = 7 if hl else None  # 7 = yellow


def remove(el) -> None:
    el.getparent().remove(el)


def has_break_or_section(el) -> bool:
    if el.find(".//" + qn("w:sectPr")) is not None:
        return True
    return any(b.get(qn("w:type")) == "page" for b in el.iter(qn("w:br")))


# ----------------------------------------------------------- body writer

class Writer:
    """Appends formatted blocks before a fixed anchor element."""

    def __init__(self, doc, anchor):
        self.doc = doc
        self.anchor = anchor
        self.fig_no: dict[int, int] = {}
        self.tab_no: dict[int, int] = {}
        self.chapter = 0
        self.figures: list[str] = []
        self.tables: list[str] = []
        self.toc: list[tuple[int, str]] = []   # (level, text)

    def _p(self) -> Paragraph:
        el = OxmlElement("w:p")
        self.anchor.addprevious(el)
        return Paragraph(el, self.doc._body)

    def para(self, text="", *, size=12, bold=False, italic=False, align="justify",
             spacing=1.15, before=0, after=6, keep_next=False, page_break=False,
             indent_first=0.0, left=0.0, hanging=0.0, font=FONT) -> Paragraph:
        p = self._p()
        pf = p.paragraph_format
        pf.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                        "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
        pf.line_spacing = spacing
        pf.space_before, pf.space_after = Pt(before), Pt(after)
        pf.keep_with_next = keep_next
        pf.page_break_before = page_break
        if indent_first:
            pf.first_line_indent = Inches(indent_first)
        if left:
            pf.left_indent = Inches(left)
        if hanging:
            pf.left_indent = Inches(left + hanging)
            pf.first_line_indent = Inches(-hanging)
        if text:
            _font(p.add_run(text), size, bold, italic, font)
        return p

    # headings -----------------------------------------------------------
    def chapter_heading(self, title: str, numbered: bool = True) -> None:
        if numbered:
            self.chapter += 1
            self.para(f"CHAPTER {self.chapter}", size=16, bold=True, align="center",
                      spacing=1.5, after=0, keep_next=True, page_break=True)
            self.para(title, size=16, bold=True, align="center", spacing=1.5, after=12,
                      keep_next=True)
            self.toc.append((0, f"{self.chapter}.|{title}"))
        else:
            self.para(title, size=14, bold=True, align="center", spacing=1.5, after=12,
                      keep_next=True, page_break=True)
            self.toc.append((0, f"|{title}"))

    def h2(self, text: str) -> None:
        self.para(text, size=14, bold=True, align="left", spacing=1.5, before=10, after=4,
                  keep_next=True)
        self.toc.append((1, text))

    def h3(self, text: str, toc: bool) -> None:
        self.para(text, size=12, bold=True, align="left", spacing=1.5, before=6, after=2,
                  keep_next=True)
        if toc:
            self.toc.append((2, text))

    # content ------------------------------------------------------------
    def body(self, text: str) -> None:
        self.para(text)

    def lead(self, lead: str, text: str) -> None:
        p = self.para()
        _font(p.add_run(lead), 12, True)
        _font(p.add_run(text), 12)

    def bullets(self, items: list[str]) -> None:
        for it in items:
            p = self.para(left=0.25, hanging=0.25, after=3)
            _font(p.add_run("•\t" + it), 12)

    def caption(self, text: str, above: bool) -> None:
        label, _, rest = text.partition(" ")
        num, _, title = rest.partition(" ")
        p = self.para(align="center", size=10, spacing=1.0, before=2 if above else 4,
                      after=4 if above else 10, keep_next=above)
        _font(p.add_run(f"{label} {num} "), 10, True)
        _font(p.add_run(title), 10)

    def _number(self, kind: str, given: str) -> str:
        counter = self.fig_no if kind == "Figure" else self.tab_no
        counter[self.chapter] = counter.get(self.chapter, 0) + 1
        num = f"{self.chapter}.{counter[self.chapter]}" if self.chapter else \
            f"A.{counter[self.chapter]}"
        title = re.sub(r"^(Figure|Table|Listing)\s+[\dA.]+\s*", "", given)
        return f"{kind} {num} {title}"

    def figure(self, file: str, caption: str, width: float) -> None:
        p = self.para(align="center", spacing=1.0, before=6, after=0, keep_next=True)
        p.add_run().add_picture(str(C.FIG / file), width=Inches(min(width, TEXT_WIDTH_IN)))
        cap = self._number("Figure", caption)
        self.caption(cap, above=False)
        self.figures.append(cap)

    def table(self, caption: str, header: list[str], rows: list[list[str]],
              widths: list[float]) -> None:
        cap = self._number("Table", caption)
        self.caption(cap, above=True)
        self.tables.append(cap)
        scale = min(1.0, TEXT_WIDTH_IN / sum(widths))
        widths = [w * scale for w in widths]
        size = 9 if len(header) >= 5 else 10
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        self.anchor.addprevious(t._tbl)
        for r_i, values in enumerate([header] + rows):
            row = t.rows[r_i]
            if r_i == 0:
                trpr = row._tr.get_or_add_trPr()
                hdr = OxmlElement("w:tblHeader")
                trpr.append(hdr)
            for c_i, v in enumerate(values):
                cell = row.cells[c_i]
                cell.width = Inches(widths[c_i])
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.line_spacing = 1.0
                _font(p.add_run(str(v)), size, bold=(r_i == 0))
                if r_i == 0:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), "D9E2F3")
                    cell._tc.get_or_add_tcPr().append(shd)
        self.para(after=4, spacing=1.0)

    def code(self, text: str, caption: str | None) -> None:
        lines = text.splitlines()
        for i, line in enumerate(lines):
            p = self.para(align="left", spacing=1.0, after=0, before=4 if i == 0 else 0,
                          keep_next=i < len(lines) - 1 or bool(caption))
            ppr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "F2F2F2")
            ppr.append(shd)
            run = p.add_run(line if line else " ")
            _font(run, 8.5, name="Courier New")
            run.font.color.rgb = RGBColor(0x1F, 0x1F, 0x1F)
        if caption:
            self.caption(caption, above=False)
        else:
            self.para(after=4, spacing=1.0)

    def block(self, b) -> None:
        kind = b[0]
        if kind == "h2":
            self.h2(b[1])
        elif kind == "h3":
            self.h3(b[1], toc=self.chapter in (3, 4))
        elif kind == "p":
            self.body(b[1])
        elif kind == "lead":
            self.lead(b[1], b[2])
        elif kind == "bullets":
            self.bullets(b[1])
        elif kind == "fig":
            self.figure(b[1], b[2], b[3])
        elif kind == "table":
            self.table(b[1], b[2], b[3], b[4])
        elif kind == "code":
            self.code(b[1], b[2])
        else:
            raise ValueError(kind)


# ------------------------------------------------------------- front matter

def roman(n: int) -> str:
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v:
            out += s
            n -= v
    return out


def find(doc, starts: str) -> Paragraph:
    for p in doc.paragraphs:
        if p.text.strip().startswith(starts):
            return p
    raise KeyError(starts)


def fill_front_matter(doc) -> None:
    title = C.TITLE
    set_text(find(doc, "<Title of the Project>"), title)
    remove(find(doc, "(Times New Roman 18")._p)
    set_text(find(doc, "Programme"), "Computer Science and Engineering")
    remove(find(doc, "(mention names of team members")._p)

    t = doc.tables[0]
    for i, (reg, name) in enumerate(C.STUDENTS):
        set_text(t.rows[i].cells[0].paragraphs[0], reg)
        set_text(t.rows[i].cells[1].paragraphs[0], name)
    remove(t.rows[2]._tr)

    g = doc.tables[1]
    set_text(g.rows[0].cells[0].paragraphs[0], C.GUIDE)
    set_text(g.rows[1].cells[0].paragraphs[0], C.DESIGNATION)

    # Declaration: written for a two-person team.
    p = find(doc, "I    hereby")
    set_runs(p, [(f"We hereby declare that the project entitled {C.TITLE_CASE} submitted by us, "
                  "for the award of the degree of Bachelor of Technology in Computer Science and "
                  "Engineering to VIT is a record of bonafide work carried out by us under the "
                  f"supervision of {C.GUIDE}.", False)])
    p = find(doc, "I further declare")
    set_runs(p, [("We further declare that the work reported in this project has not been "
                  "submitted and will not be submitted, either in part or in full, for the award "
                  "of any other degree or diploma in this institute or any other institute or "
                  "university.", False)])
    set_text(find(doc, "Signature of the Candidate"), "Signature of the Candidates")
    for note in ("Note: Include your project title and guide", "Note: Include your project title and student"):
        remove(find(doc, note)._p)

    # Certificate.
    students = " and ".join(f"{n.title()} ({r})" for r, n in C.STUDENTS)
    p = find(doc, "This is to certify")
    set_runs(p, [(f"This is to certify that the project entitled {C.TITLE_CASE} submitted by "
                  f"{students}, School of Computer Science and Engineering, VIT, for the award of "
                  "the degree of Bachelor of Technology in Computer Science and Engineering, is a "
                  "record of bonafide work carried out by them under my supervision during Fall "
                  "Semester 2026-2027, as per the VIT code of academic and research ethics.",
                  False)])
    # The longer certificate text pushes the HoD lines onto a page of their own;
    # the template's blank spacers above them absorb the difference.
    el = find(doc, "Examiner")._p.getnext()
    dropped = 0
    while el is not None and dropped < 4:
        nxt = el.getnext()
        text = "".join(t.text or "" for t in el.iter(qn("w:t"))).strip()
        if text.startswith("<<"):
            break
        if el.tag == qn("w:p") and not text and not has_break_or_section(el):
            remove(el)
            dropped += 1
        el = nxt
    # The Head of Department's name is not known to us; it stays highlighted.
    set_runs(find(doc, "<< Head of the"), [("<< Head of the Department >>", True)])
    set_text(find(doc, "<< Programme"), "B.Tech. Computer Science and Engineering")

    # Acknowledgements, in the team's voice.
    set_text(find(doc, "I am deeply grateful"),
             "We are deeply grateful to the management of Vellore Institute of Technology (VIT) "
             "for providing us with the opportunity and resources to undertake this project. "
             "Their commitment to fostering a conducive learning environment has been "
             "instrumental in our academic journey, and the support and infrastructure provided "
             "by VIT have enabled us to explore and develop our ideas to their fullest potential.")
    set_text(find(doc, "My sincere thanks to Dr. Jaisankar"),
             "Our sincere thanks to Dr. Jaisankar N, the Dean of the School of Computer Science "
             "and Engineering (SCOPE), for his unwavering support and encouragement. His "
             "leadership and vision have greatly inspired us to strive for excellence.")
    set_runs(find(doc, "I express my profound"),
             [("We express our profound appreciation to ", False),
              ("<Head of Department's Name>", True), (", the Head of the ", False),
              ("<Department Name>", True),
              (", for the insightful guidance and continuous support extended to us. The "
               "Department's commitment to a collaborative and supportive atmosphere has greatly "
               "enhanced our learning experience.", False)])
    set_text(find(doc, "I am immensely thankful"),
             f"We are immensely thankful to our project guide, {C.GUIDE}, for dedicated "
             "mentorship and invaluable feedback. The guidance we received was pivotal in "
             "refining the evaluation design of this project - in particular the decision to "
             "pilot the study before running it in full - and has enriched our overall "
             "academic experience.")
    set_text(find(doc, "Name of the Candidate"),
             "Yashika Gupta (23BCE0225)          Pratyaksh Gupta (23BCE0316)")


def fill_lists(doc, w: Writer, pages: dict[str, str]) -> None:
    # Table of contents.
    toc = doc.tables[2]
    proto_ch = copy.deepcopy(toc.rows[2]._tr)
    proto_sec = copy.deepcopy(toc.rows[3]._tr)
    proto_sub = proto_sec
    instr = next((p for p in doc.paragraphs
                  if p.text.strip().startswith("<Contents, Times New Roman 12, Line spacing 1.5>")),
                 None)
    if instr is not None:
        remove(instr._p)
    proto_abs = copy.deepcopy(toc.rows[1]._tr)
    for r in list(toc.rows)[1:]:
        remove(r._tr)

    def add(proto, a, b, c):
        el = copy.deepcopy(proto)
        toc._tbl.append(el)
        from docx.table import _Row
        row = _Row(el, toc)
        for cell, v in zip(row.cells, (a, b, c)):
            set_text(cell.paragraphs[0], v)

    add(proto_abs, "", "Abstract", pages.get("ABSTRACT", ""))
    for level, text in w.toc:
        if level == 0:
            num, _, title = text.partition("|")
            add(proto_ch, num, title, pages.get(title, ""))
        elif level == 1:
            add(proto_sec, "", text, pages.get(text, ""))
        else:
            add(proto_sub, "", "     " + text, pages.get(text, ""))

    # List of figures.
    lof = doc.tables[3]
    proto = copy.deepcopy(lof.rows[1]._tr)
    for r in list(lof.rows)[1:]:
        remove(r._tr)
    for cap in w.figures:
        _, num, title = cap.split(" ", 2)
        el = copy.deepcopy(proto)
        lof._tbl.append(el)
        from docx.table import _Row
        row = _Row(el, lof)
        for cell, v in zip(row.cells, (num, title, pages.get(cap, ""))):
            set_text(cell.paragraphs[0], v)
    remove(find(doc, "(In the chapters, figure caption")._p)

    # List of tables: tab-separated paragraphs, one per table.
    tpl = find(doc, "2.1\tTable caption") if any(
        p.text.startswith("2.1\tTable caption") for p in doc.paragraphs) else None
    if tpl is None:
        tpl = next(p for p in doc.paragraphs if p.text.startswith("2.1") and "Table caption" in p.text)
    anchor = tpl._p
    for cap in w.tables:
        _, num, title = cap.split(" ", 2)
        el = copy.deepcopy(tpl._p)
        anchor.addnext(el)
        anchor = el
        para = Paragraph(el, tpl._parent)
        runs = para.runs
        runs[0].text = num
        runs[1].text = "\t" + title
        runs[2].text = "\t" + pages.get(cap, "")
    remove(tpl._p)

    # Abbreviations, alphabetical, replacing the template's examples.
    first = find(doc, "2G")
    examples = [p for p in doc.paragraphs
                if re.match(r"^(2G|3GPP|3G|4G|AWGN|BBE|CSMA|DAS)\t", p.text)]
    anchor = first._p
    for abbr, meaning in C.ABBREVIATIONS:
        el = copy.deepcopy(first._p)
        anchor.addnext(el)
        anchor = el
        para = Paragraph(el, first._parent)
        set_text(para, f"{abbr}\t{meaning}")
    for p in examples:
        remove(p._p)
    remove(find(doc, "Note: All the abbreviations")._p)

    # Symbols and notations.
    sym_paras = [p for p in doc.paragraphs if p.text.endswith(("CFO", "NCFO"))]
    anchor = sym_paras[-1]._p
    for sym, meaning in C.SYMBOLS:
        el = copy.deepcopy(sym_paras[0]._p)
        anchor.addnext(el)
        anchor = el
        para = Paragraph(el, sym_paras[0]._parent)
        for r in para.runs[1:]:
            r._r.getparent().remove(r._r)
        para.runs[0].text = f"{sym}\t{meaning}"
        _font(para.runs[0], 12, italic=False)
        pf = para.paragraph_format
        pf.left_indent, pf.first_line_indent = Inches(1.5), Inches(-1.5)
        pf.tab_stops.add_tab_stop(Inches(1.5))
    for p in sym_paras:
        remove(p._p)


def fill_abstract(doc) -> None:
    head = find(doc, "ABSTRACT")
    # The template's page break sat inside its last symbol paragraph, which is
    # replaced; the heading carries its own break instead.
    head.paragraph_format.page_break_before = True
    head.paragraph_format.space_after = Pt(12)
    instr = [find(doc, "(Times New Roman 16"), find(doc, "<Contents, Times New Roman 12, Line spacing 1.15>"),
             find(doc, "One page and not exceeding")]
    anchor = instr[1]._p
    proto = instr[1]
    for i, text in enumerate(C.ABSTRACT + [None]):
        el = copy.deepcopy(proto._p)
        anchor.addprevious(el)
        para = Paragraph(el, proto._parent)
        if text is None:
            para.runs[0].text = ""
            for r in para.runs[1:]:
                remove(r._r)
            _font(para.runs[0], 12, True)
            para.runs[0].text = "Keywords: "
            run = para.add_run(C.KEYWORDS)
            _font(run, 12, False)
        else:
            set_text(para, text)
        pf = para.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.line_spacing = 1.15
        pf.space_after = Pt(6)
        for r in para.runs:
            _font(r, 11.5)
    for p in instr:
        remove(p._p)
    # Blank paragraphs above the heading would start the abstract mid-page.
    el = head._p.getprevious()
    while el is not None and el.tag == qn("w:p") and not has_break_or_section(el) \
            and not "".join(t.text or "" for t in el.iter(qn("w:t"))).strip():
        prev = el.getprevious()
        remove(el)
        el = prev
    # Blank spacer paragraphs between the abstract and the section break would
    # push the abstract past one page; the break itself is kept.
    el = head._p.getnext()
    while el is not None and el.find(".//" + qn("w:sectPr")) is None:
        nxt = el.getnext()
        if el.tag == qn("w:p") and not "".join(t.text or "" for t in el.iter(qn("w:t"))).strip() \
                and not has_break_or_section(el) and el.find(".//" + qn("w:drawing")) is None:
            remove(el)
        el = nxt


# ------------------------------------------------------------------- build

def build(template: Path, out: Path, pages: dict[str, str]) -> Writer:
    doc = docx.Document(template)
    fill_front_matter(doc)
    fill_abstract(doc)

    body = doc.element.body
    ch1 = find(doc, "CHAPTER 1")._p
    # The paragraph carrying the chapter section's properties (arabic from 1).
    el = ch1
    while el.find(".//" + qn("w:sectPr")) is None:
        el = el.getnext()
    chapter_section_end = el
    # Clear the template's placeholder chapter and everything after it, keeping
    # the two section-property carriers.
    el = ch1
    while el is not None and el.tag != qn("w:sectPr"):
        nxt = el.getnext()
        if el is not chapter_section_end:
            remove(el)
        el = nxt
    final_sect = body.find(qn("w:sectPr"))

    w = Writer(doc, chapter_section_end)
    for i, (title, blocks) in enumerate(C.chapters()):
        w.chapter_heading(title)
        for b in blocks:
            w.block(b)
        if i == 0:
            # The first chapter heading must not force a blank page after the
            # section break that already starts it.
            pass

    # References and the appendix go in the final section.
    w.anchor = final_sect
    w.chapter_heading("REFERENCES")
    for i, ref in enumerate(C.REFERENCES, 1):
        p = w.para(f"[{i}]\t{ref}", hanging=0.4, after=4)
    w.chapter = 0
    w.chapter_heading("APPENDIX A – SAMPLE CODE", numbered=False)
    w.body("The excerpts below are taken verbatim from the repository. They are the parts of "
           "the implementation on which the validity of the evaluation most directly depends.")
    for title, code in C.appendix_snippets():
        w.h3(title, toc=False)
        w.code(code, None)

    # The first CHAPTER heading already follows a section break.
    first = find(doc, "CHAPTER 1")
    first.paragraph_format.page_break_before = False

    fill_lists(doc, w, pages)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return w


# ---------------------------------------------------------- page numbers

def export_pdf(docx_path: Path) -> Path | None:
    if sys.platform == "win32":
        from office_pdf import export_pdf_windows

        return export_pdf_windows(docx_path)
    pdf = docx_path.with_suffix(".pdf")
    script = f'''
    tell application "Microsoft Word"
      open POSIX file "{docx_path}"
      delay 2
      set d to active document
      save as d file name "{pdf}" file format format PDF
      try
        close every document saving no
      end try
    end tell'''
    subprocess.run(["osascript", "-e", script], check=True, capture_output=True)
    return pdf


def locate_pages(pdf: Path, w: Writer) -> dict[str, str]:
    import pymupdf

    doc = pymupdf.open(pdf)
    lines = [[ln.strip() for ln in page.get_text().splitlines() if ln.strip()] for page in doc]
    decl = next(i for i, ls in enumerate(lines) if "DECLARATION" in ls)
    ch1 = next(i for i, ls in enumerate(lines) if "CHAPTER 1" in ls)
    pages: dict[str, str] = {}
    abstract = next(i for i, ls in enumerate(lines) if "ABSTRACT" in ls and i > decl + 3)
    pages["ABSTRACT"] = roman(abstract - decl + 1)

    def first_page(pred) -> str:
        for i in range(ch1, len(lines)):
            if any(pred(ln) for ln in lines[i]):
                return str(i - ch1 + 1)
        return ""

    for level, text in w.toc:
        key = text.partition("|")[2] if level == 0 else text
        pages[key] = first_page(lambda ln, k=key: ln == k)
    for cap in w.figures + w.tables:
        label = " ".join(cap.split(" ", 2)[:2])
        pages[cap] = first_page(lambda ln, k=label: ln.startswith(k + " ") or ln == k)
    return pages


def main(argv=None) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--template", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--no-pdf", action="store_true", help="skip the Word page-number pass")
    ap.add_argument("--pages-from", metavar="PDF",
                    help="take page numbers from this PDF of the --no-pdf build, saved from "
                         "any Word (desktop or web), instead of driving Word here")
    args = ap.parse_args(argv)
    template, out = Path(args.template).expanduser(), Path(args.out).expanduser()

    w = build(template, out, {})
    if args.pages_from:
        # Filling the page column does not reflow the document, so one pass is enough.
        pages = locate_pages(Path(args.pages_from).expanduser(), w)
        missing = [k for k, v in pages.items() if not v]
        build(template, out, pages)
        print("wrote", out, "(page numbers from", args.pages_from + ")")
        if missing:
            print("could not locate:", missing)
        return 0
    if args.no_pdf:
        print("wrote", out, "(no page numbers)")
        return 0
    first_pdf = export_pdf(out)
    if first_pdf is None:
        print("wrote", out, "(no page numbers: no PDF export available)")
        return 0
    pages = locate_pages(first_pdf, w)
    missing = [k for k, v in pages.items() if not v]
    w = build(template, out, pages)
    pdf = export_pdf(out)
    print("wrote", out)
    if pdf:
        print("wrote", pdf)
    if missing:
        print("could not locate:", missing)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
