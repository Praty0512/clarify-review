"""Dump every piece of visible text from a .docx or .pptx, in document order.

Used to diff a rebuilt deliverable against an earlier one, and to cross-check
the numbers in a build against paper/numbers.json.

    python docs/dump_text.py file.docx > file.txt
    python docs/dump_text.py a.docx b.docx   # unified diff of the two dumps
"""

from __future__ import annotations

import difflib
import sys
from pathlib import Path


def dump_docx(path: Path) -> list[str]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    out: list[str] = []

    def walk(parent_el, parent):
        for el in parent_el.iterchildren():
            tag = el.tag.rsplit("}", 1)[-1]
            if tag == "p":
                t = Paragraph(el, parent).text
                if t.strip():
                    out.append(t)
            elif tag == "tbl":
                tbl = Table(el, parent)
                for ri, row in enumerate(tbl.rows):
                    cells, seen = [], set()
                    for c in row.cells:
                        if id(c._tc) in seen:
                            continue
                        seen.add(id(c._tc))
                        cells.append(c.text.replace("\n", " / "))
                    out.append(f"[T r{ri}] " + " | ".join(cells))
            elif tag == "sdt":
                content = el.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}sdtContent")
                if content is not None:
                    walk(content, parent)

    walk(d.element.body, d)
    for si, sec in enumerate(d.sections):
        for kind, part in (("header", sec.header), ("footer", sec.footer)):
            for p in part.paragraphs:
                if p.text.strip():
                    out.append(f"[S{si} {kind}] {p.text}")
    return out


def _shape_text(shape, out: list[str], prefix: str) -> None:
    if shape.shape_type == 6:  # group
        for s in shape.shapes:
            _shape_text(s, out, prefix)
        return
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for p in shape.text_frame.paragraphs:
            t = "".join(r.text for r in p.runs)
            if t.strip():
                out.append(f"{prefix}[{shape.name}] {t}")
    if getattr(shape, "has_table", False) and shape.has_table:
        for ri, row in enumerate(shape.table.rows):
            out.append(f"{prefix}[{shape.name} r{ri}] " + " | ".join(c.text.replace("\n", " / ") for c in row.cells))
    if shape.shape_type == 13:  # picture
        try:
            out.append(f"{prefix}[{shape.name}] <picture {shape.image.sha1[:10]} {len(shape.image.blob)}B>")
        except Exception:
            pass


def dump_pptx(path: Path) -> list[str]:
    from pptx import Presentation

    prs = Presentation(str(path))
    out: list[str] = []
    for i, slide in enumerate(prs.slides, 1):
        pre = f"S{i:02d} "
        for shape in slide.shapes:
            _shape_text(shape, out, pre)
        if slide.has_notes_slide:
            t = slide.notes_slide.notes_text_frame.text
            if t.strip():
                out.append(f"{pre}[notes] {t}")
    return out


def dump(path: Path) -> list[str]:
    path = Path(path)
    if path.suffix.lower() == ".docx":
        return dump_docx(path)
    if path.suffix.lower() == ".pptx":
        return dump_pptx(path)
    raise SystemExit(f"unsupported: {path}")


def main(argv=None) -> int:
    args = list(sys.argv[1:] if argv is None else argv)
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", newline="\n")
    if len(args) == 1:
        print("\n".join(dump(Path(args[0]))))
        return 0
    if len(args) == 2:
        a, b = dump(Path(args[0])), dump(Path(args[1]))
        diff = list(difflib.unified_diff(a, b, args[0], args[1], lineterm="", n=0))
        print("\n".join(diff))
        return 1 if diff else 0
    raise SystemExit(__doc__)


if __name__ == "__main__":
    raise SystemExit(main())
